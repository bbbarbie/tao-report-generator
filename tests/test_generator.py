"""The generated workbook must be usable in place of the hand-made one."""

from __future__ import annotations

import openpyxl
import pytest

from app.generator import write_workbook
from tests.conftest import GROUND_TRUTH


@pytest.fixture(scope="module")
def generated(july_report, tmp_path_factory):
    path = tmp_path_factory.mktemp("out") / "TAO Compare 202607.xlsx"
    write_workbook(july_report, path)
    return openpyxl.load_workbook(path)


@pytest.fixture(scope="module")
def expected():
    if not GROUND_TRUTH.exists():
        pytest.skip("no ground-truth workbook")
    return openpyxl.load_workbook(GROUND_TRUTH)


def test_the_report_sheet_is_present_alongside_review_and_audit(generated):
    assert generated.sheetnames == ["CNTAO", "Review", "Audit"]


def test_headers_match_the_historical_workbook(generated, expected):
    ours, theirs = generated["CNTAO"], expected["CNTAO"]
    for col in range(2, 12):
        assert ours.cell(2, col).value == theirs.cell(2, col).value


def test_the_table_occupies_the_same_cells(generated, expected):
    assert generated["CNTAO"].dimensions == expected["CNTAO"].dimensions


def test_number_formats_match(generated, expected):
    ours, theirs = generated["CNTAO"], expected["CNTAO"]
    for col in (6, 7, 10, 11):  # ATA, ATB, W/B, average
        assert ours.cell(3, col).number_format == theirs.cell(3, col).number_format


def test_cells_are_bordered_and_centred_like_the_original(generated):
    ws = generated["CNTAO"]
    for row in range(2, 30):
        for col in range(2, 12):
            cell = ws.cell(row, col)
            assert cell.border.left.style == "thin"
            assert cell.alignment.vertical == "center"


def test_column_widths_and_gridlines_match(generated, expected):
    ours, theirs = generated["CNTAO"], expected["CNTAO"]
    assert ours.sheet_view.showGridLines == theirs.sheet_view.showGridLines
    for letter in ("B", "C", "D", "F", "G", "H", "I", "J", "K"):
        assert (
            round(ours.column_dimensions[letter].width, 3)
            == round(theirs.column_dimensions[letter].width, 3)
        )


def test_values_are_written_as_numbers_and_dates_not_text(generated):
    ws = generated["CNTAO"]
    from datetime import datetime

    assert isinstance(ws.cell(3, 6).value, datetime)
    assert isinstance(ws.cell(3, 8).value, (int, float))


def test_no_formulas_are_written(generated):
    ws = generated["CNTAO"]
    for row in ws.iter_rows(min_row=2, max_row=29, min_col=2, max_col=11):
        for cell in row:
            assert not (isinstance(cell.value, str) and cell.value.startswith("="))


def test_the_review_sheet_states_the_three_row_counts(generated, july_report):
    ws = generated["Review"]
    text = "\n".join(
        str(c.value) for row in ws.iter_rows() for c in row if c.value is not None
    )
    for state in ("Automatic", "Reviewed", "Unresolved"):
        assert state in text
    counts = {c.value for row in ws.iter_rows(min_col=1, max_col=1) for c in row}
    for state in ("auto", "reviewed", "unresolved"):
        assert len(july_report.rows_by_resolution(state)) in counts


def test_the_review_sheet_carries_every_open_question(generated, july_report):
    ws = generated["Review"]
    text = "\n".join(
        str(c.value) for row in ws.iter_rows() for c in row if c.value is not None
    )
    assert f"Open questions ({len(july_report.issues)})" in text
    for issue in july_report.issues:
        assert issue.question in text
        # the options and the reason for the recommendation travel with it
        for option in issue.options:
            assert option.label in text


def test_the_review_sheet_still_lists_the_informational_notes(generated, july_report):
    ws = generated["Review"]
    text = "\n".join(
        str(c.value) for row in ws.iter_rows() for c in row if c.value is not None
    )
    assert f"Notes ({len(july_report.review)})" in text


def test_the_audit_sheet_records_provenance_for_every_row(generated, july_report):
    ws = generated["Audit"]
    assert ws.max_row - 1 == len(july_report.rows)
    headers = [c.value for c in ws[1]]
    assert "Arr formula" in headers and "W/B source file" in headers
    method_col = headers.index("W/B method") + 1
    methods = {ws.cell(r, method_col).value for r in range(2, ws.max_row + 1)}
    assert methods <= {"copied_from_daily", "calculated", "unavailable"}


def test_a_report_with_no_rows_still_writes_a_valid_workbook(tmp_path, daily_files):
    from app.report import build_report

    empty = build_report(daily_files, 2019, 1)
    assert not empty.rows
    path = write_workbook(empty, tmp_path / "empty.xlsx")
    wb = openpyxl.load_workbook(path)
    assert wb["CNTAO"].cell(2, 2).value == "SVC"
