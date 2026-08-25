"""Proves a packaged build actually works, without pretending to test the GUI.

Run inside the frozen application:

    "TAO Report Generator.exe" --selftest

It builds a small Daily Berth Report from scratch, puts it through the whole
engine, and checks the workbook that comes out. That exercises the parts a
packaged build most often breaks — a missing openpyxl, a data file left out of
the bundle, an import that only resolved in the source tree.

It deliberately does not claim the interface works. Driving a real window
reliably in CI is another problem; whether the toolkit *loads* is checked
separately and reported as exactly that.
"""

from __future__ import annotations

import os
import sys
import tempfile
import threading
from datetime import datetime
from pathlib import Path

CHECKS: list[tuple[str, str]] = []
TRANSCRIPT: list[str] = []

# A packaged windowed build has no console on Windows, so sys.stdout may be
# None and print() would raise. Everything goes through say(), and a copy is
# left in a file the build verifier can read either way.
RESULT_FILE = "selftest-result.txt"

# Nothing here should take more than a few seconds. If it ever does, the
# process kills itself rather than being left for a build job to wait on.
WATCHDOG_SECONDS = 120


def say(line: str = "") -> None:
    TRANSCRIPT.append(line)
    stream = sys.stdout
    if stream is None:
        return
    try:
        print(line, file=stream, flush=True)
    except (OSError, ValueError, AttributeError):
        pass


def record(name: str, detail: str = "") -> None:
    CHECKS.append((name, detail))
    say(f"  ok   {name}" + (f" — {detail}" if detail else ""))


def start_watchdog(seconds: int = WATCHDOG_SECONDS) -> threading.Timer:
    """Guarantee this process ends, whatever happens inside it."""

    def give_up() -> None:  # pragma: no cover - only fires on a real stall
        say(f"\n  FAIL the self-test stalled for more than {seconds}s")
        write_transcript()
        os._exit(3)

    timer = threading.Timer(seconds, give_up)
    timer.daemon = True
    timer.start()
    return timer


def write_transcript(folder: Path | None = None) -> Path | None:
    """Leave the result somewhere readable, since stdout may go nowhere."""
    target = Path(folder) if folder else Path(sys.executable).resolve().parent
    try:
        path = target / RESULT_FILE
        path.write_text("\n".join(TRANSCRIPT) + "\n", encoding="utf-8")
        return path
    except OSError:
        return None


def build_daily_workbook(path: Path) -> None:
    """A miniature Daily Berth Report, in the real layout."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Daily Berth Report"
    ws["B1"] = "青島每日船舶動態表"
    ws["B3"] = "WK30船舶動態："
    headers = [
        "SVC", "VSL/VOY", "TFC", "TML", "ETA", "ATA", "ETB", "ATB", "ETD", "ATD",
        "Window Time (ETA-ETD)", None, "Arr Delay Hrs", "Dep Delay Hrs",
        "Waiting", "GPH", "Delay Reason",
    ]
    for offset, label in enumerate(headers):
        if label:
            ws.cell(4, 2 + offset, label)

    rows = [
        # A WAN HAI call: the Daily prints its own figures, which are copied.
        ["CI5", "WAN HAI 501/W267", "W26751", "QQCTN", None, "22/2135", "22/2230",
         "22/2236", "23/1000", "23/0910", "20/1800 - 21/0800", "MON 18 - TUE 08",
         51.6, 49.2, 1.0, 28.5, None],
        # A partner call: the columns are blank, so the figures are calculated.
        # Arrives 12 hours before its window opens and is berthed early, which
        # exercises both the negative delay and the waiting-time floor.
        ["AP2", "NYK RUMINA/E076", "E076JONR", "QQCTN", None, "21/0800", "21/1900",
         "21/1907", "22/1300", "22/2039", "21/2000 - 22/1200", "TUE 20 - WED 12",
         None, None, None, None, None],
    ]
    for r, values in enumerate(rows, start=5):
        for offset, value in enumerate(values):
            if value is not None:
                ws.cell(r, 2 + offset, value)
    wb.save(path)


def main() -> int:
    watchdog = start_watchdog()
    try:
        return _run()
    finally:
        watchdog.cancel()
        write_transcript()


def _run() -> int:
    say("TAO Report Generator — packaged self-test")
    say(f"  python {sys.version.split()[0]}")
    say(f"  frozen: {getattr(sys, 'frozen', False)}")
    say()

    try:
        from app.calculations import build_row
        from app.generator import write_workbook
        from app.operators import OperatorMapping
        from app.parser import parse_daily_workbook
        from app.report import build_report
        from app.voyage_history import build_histories

        record("engine imports")

        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            daily = folder / "Daily 20260731.xlsx"
            build_daily_workbook(daily)
            record("workbook written", daily.name)

            snapshots = parse_daily_workbook(daily)
            assert len(snapshots) == 2, f"expected 2 rows, parsed {len(snapshots)}"
            record("Daily parsed", f"{len(snapshots)} voyages")

            whl = next(s for s in snapshots if s.tfc == "W26751")
            assert whl.ata == datetime(2026, 7, 22, 21, 35), whl.ata
            assert whl.window_start == datetime(2026, 7, 20, 18, 0), whl.window_start
            record("shipping dates resolved", f"ATA {whl.ata:%Y-%m-%d %H:%M}")

            index_whl = build_histories([whl])
            house = build_row(index_whl.resolve("W26751"), OperatorMapping.load())
            assert house.opr == "WHL", house.opr
            assert house.waiting.method == "copied_from_daily", house.waiting.method
            assert not house.flags, house.flags
            record("WAN HAI figures copied from the Daily", f"W/B {house.waiting.value}")

            index = build_histories(snapshots)
            mapping = OperatorMapping.load()
            partner = build_row(index.resolve("E076JONR"), mapping)
            assert partner.opr == "ONE", partner.opr
            record("operator resolved", f"NYK RUMINA -> {partner.opr}")

            assert partner.arr_delay.value == -12.0, partner.arr_delay.value
            assert partner.dep_delay.value == 8.7, partner.dep_delay.value
            assert partner.waiting.value == 0.0, partner.waiting.value
            record(
                "figures calculated",
                f"arr {partner.arr_delay.value}, dep {partner.dep_delay.value}, "
                f"W/B {partner.waiting.value}",
            )

            report = build_report([daily], 2026, 7, mapping=mapping)
            assert report.rows, "no rows produced"
            record("report built", f"{len(report.rows)} rows")

            out = write_workbook(report, folder / "TAO Compare 202607.xlsx")
            assert out.exists() and out.stat().st_size > 0
            record("workbook generated", f"{out.stat().st_size} bytes")

            import openpyxl

            check = openpyxl.load_workbook(out)
            assert check.sheetnames == ["CNTAO", "Review", "Audit"], check.sheetnames
            assert check["CNTAO"].cell(2, 2).value == "SVC"
            record("output verified", ", ".join(check.sheetnames))

    except Exception as exc:  # noqa: BLE001 - the exit code is the contract
        import traceback

        say(f"\n  FAIL {type(exc).__name__}: {exc}")
        for line in traceback.format_exc().splitlines():
            say(f"    {line}")
        return 1

    # Reported separately and never as a pass: a toolkit that imports is not a
    # working window, and claiming otherwise would be worse than saying nothing.
    try:
        from PySide6 import QtCore

        say(f"\n  note: Qt {QtCore.__version__} loaded (interface not exercised)")
    except Exception as exc:  # noqa: BLE001
        say(f"\n  FAIL the interface toolkit did not load: {exc}")
        return 1

    say(f"\nAll {len(CHECKS)} engine checks passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
