"""Row-by-row comparison of a generated report against a known-good workbook.

Used only for regression testing. The generator itself never reads ground truth.
"""

from __future__ import annotations

import json
from dataclasses import asdict, dataclass, field
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl

from app.calculations import MonthlyRow
from app.report import MonthlyReport

# Outcome codes
EXACT_MATCH = "EXACT_MATCH"
MATCH_WITH_TOLERANCE = "MATCH_WITH_TOLERANCE"
CALCULATION_MISMATCH = "CALCULATION_MISMATCH"
FIELD_MISMATCH = "FIELD_MISMATCH"
MISSING_FROM_GENERATED = "MISSING_FROM_GENERATED"
EXTRA_IN_GENERATED = "EXTRA_IN_GENERATED"
GROUND_TRUTH_INCONSISTENT = "GROUND_TRUTH_INCONSISTENT"

NUMERIC_TOLERANCE = 0.11  # both workbooks print hours to one decimal
TIME_TOLERANCE = timedelta(minutes=1)

GT_COLUMNS = {
    "svc": 2,
    "tfc": 3,
    "opr": 4,
    "terminal": 5,
    "ata": 6,
    "atb": 7,
    "arr_delay": 8,
    "dep_delay": 9,
    "waiting": 10,
    "average_waiting": 11,
}


@dataclass
class GroundTruthRow:
    svc: str | None
    tfc: str | None
    opr: str | None
    terminal: str | None
    ata: datetime | None
    atb: datetime | None
    arr_delay: float | None
    dep_delay: float | None
    waiting: float | None
    average_waiting: float | None
    row: int


@dataclass
class FieldDiff:
    field: str
    expected: object
    actual: object


@dataclass
class RowComparison:
    tfc: str | None
    svc: str | None
    outcome: str
    diffs: list[FieldDiff] = field(default_factory=list)
    flags: list[str] = field(default_factory=list)
    note: str = ""


@dataclass
class ValidationReport:
    period: str
    total_expected: int
    total_generated: int
    comparisons: list[RowComparison]

    @property
    def counts(self) -> dict[str, int]:
        out: dict[str, int] = {}
        for c in self.comparisons:
            out[c.outcome] = out.get(c.outcome, 0) + 1
        return out

    @property
    def matched(self) -> int:
        return sum(
            1
            for c in self.comparisons
            if c.outcome in (EXACT_MATCH, MATCH_WITH_TOLERANCE)
        )

    def to_json(self) -> str:
        return json.dumps(
            {
                "period": self.period,
                "expected_rows": self.total_expected,
                "generated_rows": self.total_generated,
                "matched": self.matched,
                "counts": self.counts,
                "rows": [asdict(c) for c in self.comparisons],
            },
            ensure_ascii=False,
            indent=2,
            default=str,
        )

    def summary(self) -> str:
        lines = [
            f"TAO Compare {self.period} — validation against ground truth",
            f"  {self.total_expected} ground-truth rows, {self.total_generated} generated rows",
            f"  {self.matched} reproduced",
        ]
        for outcome, n in sorted(self.counts.items()):
            if outcome in (EXACT_MATCH, MATCH_WITH_TOLERANCE):
                continue
            lines.append(f"  {n} {outcome.lower().replace('_', ' ')}")
        lines.append("")
        for c in self.comparisons:
            if c.outcome in (EXACT_MATCH, MATCH_WITH_TOLERANCE):
                continue
            detail = "; ".join(
                f"{d.field}: expected {d.expected!r}, got {d.actual!r}" for d in c.diffs
            )
            lines.append(f"  {c.tfc or '?'} [{c.outcome}]")
            if detail:
                lines.append(f"      {detail}")
            if c.note:
                lines.append(f"      note: {c.note}")
        return "\n".join(lines)


def load_ground_truth(path: str | Path, sheet: str = "CNTAO") -> list[GroundTruthRow]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    header_row = _find_header_row(ws)
    rows: list[GroundTruthRow] = []
    for r in range(header_row + 1, ws.max_row + 1):
        if ws.cell(r, GT_COLUMNS["svc"]).value is None:
            continue
        values = {k: ws.cell(r, c).value for k, c in GT_COLUMNS.items()}
        rows.append(GroundTruthRow(row=r, **values))
    wb.close()
    return rows


def _find_header_row(ws) -> int:
    for r in range(1, min(ws.max_row, 20) + 1):
        if str(ws.cell(r, GT_COLUMNS["svc"]).value).strip() == "SVC":
            return r
    return 2


def _match_key(tfc: str | None) -> str:
    return (tfc or "").strip().upper()


def _similar_tfc(a: str, b: str) -> bool:
    """Same length, at most one character different — a likely transcription slip."""
    if len(a) != len(b):
        return False
    return sum(1 for x, y in zip(a, b) if x != y) <= 1


def compare(report: MonthlyReport, expected: list[GroundTruthRow]) -> ValidationReport:
    generated = {_match_key(r.tfc): r for r in report.rows}
    used: set[str] = set()
    comparisons: list[RowComparison] = []

    for gt in expected:
        key = _match_key(gt.tfc)
        actual = generated.get(key)
        note = ""
        if actual is None:
            # The historical workbook is typed by hand; a TFC code that differs
            # by one character from a generated row is the same call.
            candidates = [
                r
                for k, r in generated.items()
                if k not in used and _similar_tfc(k, key) and r.svc == gt.svc
            ]
            if len(candidates) == 1:
                actual = candidates[0]
                note = f"TFC differs: ground truth {gt.tfc!r}, Daily files say {actual.tfc!r}"
        if actual is None:
            comparisons.append(
                RowComparison(gt.tfc, gt.svc, MISSING_FROM_GENERATED, note="not generated")
            )
            continue

        used.add(_match_key(actual.tfc))
        diffs = _diff(gt, actual)
        defect = _ground_truth_defect(gt)
        if diffs and defect and all(d.field in ("ata", "atb") for d in diffs):
            comparisons.append(
                RowComparison(gt.tfc, gt.svc, GROUND_TRUTH_INCONSISTENT, diffs,
                              list(actual.flags), defect)
            )
            continue
        if not diffs:
            outcome = EXACT_MATCH if not note else MATCH_WITH_TOLERANCE
        elif all(d.field in ("arr_delay", "dep_delay", "waiting", "average_waiting") for d in diffs):
            outcome = CALCULATION_MISMATCH
        else:
            outcome = FIELD_MISMATCH
        comparisons.append(
            RowComparison(gt.tfc, gt.svc, outcome, diffs, list(actual.flags), note)
        )

    for key, row in generated.items():
        if key not in used and not any(_match_key(g.tfc) == key for g in expected):
            comparisons.append(
                RowComparison(row.tfc, row.svc, EXTRA_IN_GENERATED, flags=list(row.flags))
            )

    return ValidationReport(
        period=report.period,
        total_expected=len(expected),
        total_generated=len(report.rows),
        comparisons=comparisons,
    )


def _ground_truth_defect(gt: GroundTruthRow) -> str:
    """Detect a hand-typed row that contradicts itself.

    The historical workbook is typed by a person, and a wrong month in a
    timestamp is both easy to make and easy to prove: a vessel cannot berth
    before it arrives. Only that impossibility counts as a defect — a W/B
    that looks small next to the ATA/ATB gap is usually the waiting rule
    working correctly, not a mistake.
    """
    if gt.ata is None or gt.atb is None:
        return ""
    if gt.atb < gt.ata:
        return (
            f"ground truth has ATB {gt.atb:%Y-%m-%d %H:%M} before "
            f"ATA {gt.ata:%Y-%m-%d %H:%M}, which cannot happen"
        )
    return ""


def _diff(gt: GroundTruthRow, actual: MonthlyRow) -> list[FieldDiff]:
    diffs: list[FieldDiff] = []

    def text(field_name: str, expected, got):
        if (expected or "").strip().upper() != (got or "").strip().upper():
            diffs.append(FieldDiff(field_name, expected, got))

    def number(field_name: str, expected, got):
        if expected is None and got is None:
            return
        if expected is None or got is None:
            diffs.append(FieldDiff(field_name, expected, got))
        elif abs(float(expected) - float(got)) > NUMERIC_TOLERANCE:
            diffs.append(FieldDiff(field_name, expected, got))

    def moment(field_name: str, expected, got):
        if expected is None and got is None:
            return
        if expected is None or got is None:
            diffs.append(FieldDiff(field_name, expected, got))
        elif abs(expected - got) > TIME_TOLERANCE:
            diffs.append(FieldDiff(field_name, expected, got))

    text("svc", gt.svc, actual.svc)
    text("opr", gt.opr, actual.opr)
    text("terminal", gt.terminal, actual.terminal)
    moment("ata", gt.ata, actual.ata)
    moment("atb", gt.atb, actual.atb)
    number("arr_delay", gt.arr_delay, actual.arr_delay.value)
    number("dep_delay", gt.dep_delay, actual.dep_delay.value)
    number("waiting", gt.waiting, actual.waiting.value)
    number("average_waiting", gt.average_waiting, actual.average_waiting)
    return diffs
