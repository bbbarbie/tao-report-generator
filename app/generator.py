"""Writes the TAO Compare workbook, matching the historical layout.

Layout, fonts, number formats and column widths are taken from the existing
manually-produced workbook so the generated file drops straight into the
recipient's routine. Two extra sheets — Review and Audit — carry everything
the CNTAO sheet deliberately does not show.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from app.calculations import MonthlyRow
from app.report import MonthlyReport

SHEET_NAME = "CNTAO"
REVIEW_SHEET = "Review"
AUDIT_SHEET = "Audit"

HEADERS = [
    "SVC",
    "TFC Code",
    "OPR",
    "靠泊碼頭",
    "ATA",
    "ATB",
    "Arr Delay\n(hr)",
    "Dep Delay\n(Hr)",
    "W/B (hr)",
    "平均侯泊時間",
]

FIRST_COL = 2  # data starts in column B, as in the historical workbook
HEADER_ROW = 2

COLUMN_WIDTHS = {
    "A": 5.88671875,
    "B": 6.77734375,
    "C": 9.44140625,
    "D": 8.44140625,
    "E": 9.5,
    "F": 9.88671875,
    "G": 10.0,
    "H": 6.5546875,
    "I": 5.5546875,
    "J": 5.88671875,
    "K": 11.5546875,
}

DATETIME_FORMAT = "mm/dd\\ hh:mm"
DECIMAL_FORMAT = "#,##0.0_);[Red]\\(#,##0.0\\)"

LATIN_FONT = Font(name="Calibri", size=12)
CJK_FONT = Font(name="宋体", size=12)
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style(cell, font=LATIN_FONT, align=CENTER, fmt=None):
    cell.font = font
    cell.alignment = align
    cell.border = BORDER
    if fmt:
        cell.number_format = fmt


def write_workbook(report: MonthlyReport, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.sheet_view.showGridLines = False

    for letter, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[letter].width = width
    ws.row_dimensions[HEADER_ROW].height = 48.0

    for offset, label in enumerate(HEADERS):
        col = FIRST_COL + offset
        cell = ws.cell(HEADER_ROW, col, label)
        cjk = any("一" <= ch <= "鿿" for ch in label)
        _style(cell, CJK_FONT if cjk else LATIN_FONT, CENTER_WRAP if "\n" in label else CENTER)

    for i, row in enumerate(report.rows):
        _write_row(ws, HEADER_ROW + 1 + i, row)

    _write_review(wb, report)
    _write_audit(wb, report)

    wb.save(path)
    return path


def _write_row(ws, r: int, row: MonthlyRow) -> None:
    values = [
        row.svc,
        row.tfc,
        row.opr,
        row.terminal,
        row.ata,
        row.atb,
        row.arr_delay.value,
        row.dep_delay.value,
        row.waiting.value,
        row.average_waiting,
    ]
    formats = [
        None, None, None, None,
        DATETIME_FORMAT, DATETIME_FORMAT,
        "General", "General",
        DECIMAL_FORMAT, DECIMAL_FORMAT,
    ]
    for offset, (value, fmt) in enumerate(zip(values, formats)):
        cell = ws.cell(r, FIRST_COL + offset, value)
        _style(cell, LATIN_FONT, CENTER_WRAP if offset >= 3 else CENTER, fmt)
    ws.row_dimensions[r].height = 15.75


def _write_review(wb: Workbook, report: MonthlyReport) -> None:
    ws = wb.create_sheet(REVIEW_SHEET)
    headers = ["TFC", "SVC", "Vessel", "Reason", "Detail"]
    for c, label in enumerate(headers, start=1):
        cell = ws.cell(1, c, label)
        cell.font = Font(name="Calibri", size=11, bold=True)
    for r, item in enumerate(report.review, start=2):
        for c, key in enumerate(headers, start=1):
            ws.cell(r, c, item.as_dict()[key])
    widths = [12, 7, 26, 40, 80]
    for c, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = width
    if not report.review:
        ws.cell(2, 1, "No exceptions — every row was produced from complete source data.")


def _write_audit(wb: Workbook, report: MonthlyReport) -> None:
    ws = wb.create_sheet(AUDIT_SHEET)
    headers = [
        "SVC", "TFC", "Vessel", "OPR", "OPR source", "Terminal",
        "ATA", "ATB", "ATD", "Window start", "Window end",
        "Arr Delay", "Arr method", "Arr formula", "Arr source file",
        "Dep Delay", "Dep method", "Dep formula", "Dep source file",
        "W/B", "W/B method", "W/B formula", "W/B source file",
        "Average W/B", "Flags", "Snapshots seen in",
    ]
    for c, label in enumerate(headers, start=1):
        cell = ws.cell(1, c, label)
        cell.font = Font(name="Calibri", size=11, bold=True)

    for r, row in enumerate(report.rows, start=2):
        values = [
            row.svc, row.tfc, row.vessel_voyage, row.opr, row.operator_source,
            row.terminal, row.ata, row.atb, row.atd, row.window_start, row.window_end,
            row.arr_delay.value, row.arr_delay.method, row.arr_delay.formula,
            row.arr_delay.source_file,
            row.dep_delay.value, row.dep_delay.method, row.dep_delay.formula,
            row.dep_delay.source_file,
            row.waiting.value, row.waiting.method, row.waiting.formula,
            row.waiting.source_file,
            row.average_waiting,
            ", ".join(row.flags),
            f"{len(row.snapshot_files)} files: "
            f"{row.snapshot_files[0] if row.snapshot_files else ''} .. "
            f"{row.snapshot_files[-1] if row.snapshot_files else ''}",
        ]
        for c, value in enumerate(values, start=1):
            cell = ws.cell(r, c, value)
            if hasattr(value, "year"):
                cell.number_format = "yyyy/mm/dd hh:mm"

    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 18
