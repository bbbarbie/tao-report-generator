import openpyxl, sys, glob, os
for path in ["samples/ground_truth/TAO Compare 202607.xlsx"]:
    wb = openpyxl.load_workbook(path, data_only=False)
    print("FILE:", path)
    for ws in wb.worksheets:
        print("  SHEET:", repr(ws.title), "dims", ws.dimensions, ws.max_row, ws.max_column, "state", ws.sheet_state)
        print("  merged:", ws.merged_cells.ranges)
