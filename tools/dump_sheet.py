import openpyxl, sys
path, sheet = sys.argv[1], sys.argv[2]
maxr = int(sys.argv[3]) if len(sys.argv)>3 else 10**6
wbv = openpyxl.load_workbook(path, data_only=True)
wbf = openpyxl.load_workbook(path, data_only=False)
wsv, wsf = wbv[sheet], wbf[sheet]
print("MERGED:", sorted(str(r) for r in wsf.merged_cells.ranges))
for r in range(1, min(wsf.max_row, maxr)+1):
    parts=[]
    for c in range(1, wsf.max_column+1):
        v = wsv.cell(r,c).value; f = wsf.cell(r,c).value
        if v is None and f is None: continue
        L = openpyxl.utils.get_column_letter(c)
        s=f"{L}{r}={v!r}"
        if isinstance(f,str) and f.startswith("="): s+=f"[{f}]"
        parts.append(s)
    if parts: print(" | ".join(parts))
