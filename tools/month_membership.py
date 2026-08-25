import openpyxl
from app.parser import parse_many, discover_daily_files
from app.voyage_history import build_histories
from app.normalization import is_qingdao, normalize_terminal

gt = openpyxl.load_workbook("samples/ground_truth/TAO Compare 202607.xlsx", data_only=True)["CNTAO"]
gt_rows=[]
for r in range(3, gt.max_row+1):
    if gt.cell(r,2).value is None: continue
    gt_rows.append({k: gt.cell(r,c).value for k,c in
        [("svc",2),("tfc",3),("opr",4),("tml",5),("ata",6),("atb",7),("arr",8),("dep",9),("wb",10),("avg",11)]})
print(len(gt_rows), "ground-truth rows")

snaps = parse_many(discover_daily_files("samples/daily"))
idx = build_histories(snaps)
cand=[]
for h in idx.all():
    c = h.consolidated()
    cand.append((h, c))

gt_tfc = {r["tfc"] for r in gt_rows}
def july(dt): return dt is not None and dt.year==2026 and dt.month==7

sel = [(h,c) for h,c in cand if is_qingdao(c.terminal) and july(c.atd)]
print("\nQingdao + ATD in July:", len(sel))
sel_tfc = {h.voyage_key for h,c in sel}
print("  in GT but not selected:", sorted(gt_tfc - sel_tfc))
print("  selected but not in GT:")
for h,c in sorted(sel, key=lambda x:(x[1].svc or "", x[1].atd)):
    if h.voyage_key not in gt_tfc:
        print(f"    {c.svc:4} {h.voyage_key:10} {normalize_terminal(c.terminal):7} vsl={c.vessel_voyage} ATD={c.atd} ATA={c.ata}")
print("\n--- GT rows and their match status ---")
for r in gt_rows:
    h = idx.resolve(r["tfc"])
    if h is None:
        print(f"  {r['svc']:4} {r['tfc']:10} {r['opr']:5} NO TFC MATCH in dailies (GT ATA={r['ata']} ATB={r['atb']})")
    else:
        c=h.consolidated()
        print(f"  {r['svc']:4} {r['tfc']:10} {r['opr']:5} tml={normalize_terminal(c.terminal):7} ATD={c.atd} july={july(c.atd)}")
