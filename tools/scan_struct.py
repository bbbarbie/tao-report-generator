import openpyxl, glob, os, re, collections
hdrs=collections.Counter(); wins=collections.Counter(); svcs=collections.Counter(); tmls=collections.Counter()
secmark=collections.Counter(); types=collections.Counter()
for p in sorted(glob.glob("samples/daily/*.xlsx")):
    wb=openpyxl.load_workbook(p,data_only=True); ws=wb["Daily Berth Report"]
    for r in range(1, ws.max_row+1):
        row=[ws.cell(r,c).value for c in range(1, ws.max_column+1)]
        txt=[str(v).strip() for v in row if v is not None]
        if any(re.search(r'WK\s*\d+', t) and '動態' in t for t in txt):
            secmark[tuple(t for t in txt)]+=1
        if 'SVC' in txt and 'TFC' in txt:
            hdrs[tuple((openpyxl.utils.get_column_letter(i+1), str(v).strip().replace("\n"," ")) for i,v in enumerate(row) if v is not None)]+=1
for p in sorted(glob.glob("samples/daily/*.xlsx")):
    wb=openpyxl.load_workbook(p,data_only=True); ws=wb["Daily Berth Report"]
    for r in range(1, ws.max_row+1):
        v=ws.cell(r,12).value  # L
        if v is not None and r>30:
            s=str(v).strip()
            if not re.fullmatch(r'\d{1,2}/\d{4}\s*-\s*\d{1,2}/\d{4}', s): wins[s]+=1
        b=ws.cell(r,2).value; e=ws.cell(r,5).value
        if b and r>30: svcs[str(b).strip()]+=1
        if e and r>30: tmls[str(e).strip()]+=1
        for c in (6,7,8,9,10,11):
            val=ws.cell(r,c).value
            if val is not None and r>34: types[type(val).__name__]+=1
print("SECTION MARKERS:"); [print("  ",k,v) for k,v in secmark.items()]
print("\nHEADER VARIANTS:", len(hdrs))
for k,v in hdrs.items(): print("  n=%d"%v, k)
print("\nNON-STANDARD WINDOW STRINGS:"); [print("  %r x%d"%(k,v)) for k,v in wins.most_common(40)]
print("\nSVCs:", dict(svcs))
print("\nTMLs:", dict(tmls))
print("\nTIME CELL TYPES:", dict(types))
