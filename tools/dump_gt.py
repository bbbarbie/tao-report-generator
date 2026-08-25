import openpyxl
path="samples/ground_truth/TAO Compare 202607.xlsx"
wbv = openpyxl.load_workbook(path, data_only=True)
wbf = openpyxl.load_workbook(path, data_only=False)
for name in wbf.sheetnames:
    wsv, wsf = wbv[name], wbf[name]
    print("="*100)
    print("SHEET", name, wsf.dimensions)
    for r in range(1, min(wsf.max_row,120)+1):
        vals=[]
        for c in range(1, wsf.max_column+1):
            f = wsf.cell(r,c).value
            v = wsv.cell(r,c).value
            if f is None and v is None: continue
            s = f"{openpyxl.utils.get_column_letter(c)}{r}={v!r}"
            if isinstance(f,str) and f.startswith("="): s += f" [F:{f}]"
            vals.append(s)
        if vals: print(" | ".join(vals))
